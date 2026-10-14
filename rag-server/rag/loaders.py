"""원천 자료 → chunk 변환기.

chunk 공통 필드
- id, doc_id, collection, title, year, year_from, year_to, audience, status, publisher, source_url
- location: 근거 위치(PDF 페이지, 통계표 제목, 장소명 등)
- text: LLM에 전달할 본문(토큰 절약을 위해 군더더기를 제거한 형태)
- meta: 장소 좌표·행사 기간 등 부가 정보
"""

import csv
import io
import itertools
import json
import logging
import re
from collections import defaultdict
from collections.abc import Iterator
from datetime import date, timedelta
from pathlib import Path

from rag.manifest import SourceDocument
from rag.text import clean_html, compact, one_line

logger = logging.getLogger(__name__)

PDF_CHUNK_CHARS = 900
PDF_CHUNK_OVERLAP = 150
TABLE_CHUNK_CHARS = 1400
MIN_PAGE_CHARS = 40


# ---------------------------------------------------------------------------
# 공통 유틸
# ---------------------------------------------------------------------------

def _chunk(document: SourceDocument, seq: int, location: str, text: str, meta: dict | None = None) -> dict:
    return {
        "id": f"{document.id}:{seq:05d}",
        **document.base_meta(),
        "location": location,
        "text": text,
        "meta": meta or {},
    }


def _split_text(text: str, size: int, overlap: int) -> list[str]:
    """줄 단위로 누적하다가 size를 넘으면 자른다. 다음 조각은 overlap만큼 앞 내용을 이어 받는다."""
    lines = text.split("\n")
    pieces: list[str] = []
    buffer = ""
    for line in lines:
        while len(line) > size:  # 아주 긴 한 줄은 강제로 나눈다.
            head, line = line[:size], line[size:]
            if buffer:
                pieces.append(buffer)
                buffer = ""
            pieces.append(head)
        candidate = f"{buffer}\n{line}" if buffer else line
        if len(candidate) > size and buffer:
            pieces.append(buffer)
            tail = buffer[-overlap:] if overlap else ""
            buffer = f"{tail}\n{line}" if tail else line
        else:
            buffer = candidate
    if buffer.strip():
        pieces.append(buffer)
    return [piece.strip() for piece in pieces if piece.strip()]


def _read_csv(path: Path) -> list[dict[str, str]]:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "cp949"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValueError(f"CSV 인코딩을 판별하지 못했습니다: {path.name}")
    return [
        {(key or "").strip(): (value or "").strip() for key, value in row.items()}
        for row in csv.DictReader(io.StringIO(text))
    ]


def _fmt_number(value) -> str:
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return f"{value:.1f}"
    return str(value)


def _parse_period(value: str | None) -> tuple[str | None, str | None]:
    """'2026-12-24~2026-12-31' → ('2026-12-24', '2026-12-31')"""
    if not value:
        return None, None
    found = re.findall(r"\d{4}-\d{2}-\d{2}", value)
    if not found:
        return None, None
    return found[0], found[-1]


# ---------------------------------------------------------------------------
# PDF 보고서
# ---------------------------------------------------------------------------

RUNNING_HEADER_PATTERNS = (
    re.compile(r"^\[(?:내국인|외국인)\]\s*\d{4}년 서울관광 실태조사\s*\d*$"),
    re.compile(r"^\d{4} 서울관광 실태조사 결과보고서\((?:내국인|외국인)\)$"),
    re.compile(r"^[-–\s\d|ㅣ]+$"),  # 쪽번호
    re.compile(r"^[가-힣]$"),  # 세로 탭 글자(조/사/개/요 …)
)
QUESTION_CODE_PATTERN = re.compile(r"(?m)^(?:[A-Z]{1,3}\d+(?:[-_]\d+)*[\].)]|Part [A-Z]\.)")
CIRCLED_NUMBERS = set("①②③④⑤⑥⑦⑧⑨⑩⑪⑫⑬⑭⑮")
OVERVIEW_PATTERN = re.compile(r"주요 용어 정리|조사 목적|표본\s?설계|표본배분|조사 개요|조사개요")


def _strip_page_noise(text: str) -> str:
    """보고서 머리말·쪽번호·세로 탭 글자를 제거해 토큰을 줄인다."""
    return "\n".join(
        line for line in text.split("\n")
        if not any(pattern.match(line) for pattern in RUNNING_HEADER_PATTERNS)
    )


def _looks_like_questionnaire(text: str) -> bool:
    """보고서 부록의 설문지 페이지는 결과 수치가 없어 검색 노이즈가 되므로 제외한다."""
    circled = sum(1 for char in text if char in CIRCLED_NUMBERS)
    return len(QUESTION_CODE_PATTERN.findall(text)) >= 2 or circled >= 10


def load_pdf(document: SourceDocument, path: Path) -> Iterator[dict]:
    from pypdf import PdfReader

    logging.getLogger("pypdf").setLevel(logging.ERROR)
    reader = PdfReader(str(path))
    seq = 0
    for page_number, page in enumerate(reader.pages, start=1):
        try:
            text = compact(page.extract_text() or "")
        except Exception as error:  # 손상된 폰트 객체 등 페이지 단위 오류는 건너뛴다.
            logger.warning("PDF 페이지 추출 실패 %s p.%s: %r", path.name, page_number, error)
            continue
        text = _strip_page_noise(text)
        if len(text) < MIN_PAGE_CHARS or _looks_like_questionnaire(text):
            continue
        meta = {"page": page_number}
        if OVERVIEW_PATTERN.search(text):
            meta["section"] = "overview"  # 조사 개요·용어 정의 페이지(검색 시 가중치를 낮춘다)
        for piece in _split_text(text, PDF_CHUNK_CHARS, PDF_CHUNK_OVERLAP):
            yield _chunk(document, seq, f"PDF p.{page_number}", piece, meta)
            seq += 1


# ---------------------------------------------------------------------------
# 통계표 XLSX
# ---------------------------------------------------------------------------

def _is_number(value) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def _column_labels(header_rows: list[tuple], width: int) -> list[str]:
    labels = []
    for column in range(width):
        parts = []
        for row in header_rows:
            value = row[column] if column < len(row) else None
            if value not in (None, ""):
                parts.append(one_line(str(value)))
        label = " ".join(parts)
        label = re.sub(r"\s*%$", "", label).strip()
        labels.append(label)
    return labels


def _render_table(title: str, rows: list[tuple], note: str) -> list[str]:
    header_rows: list[tuple] = []
    data_rows: list[tuple] = []
    for row in rows:
        if any(_is_number(value) for value in row):
            data_rows.append(row)
        elif not data_rows:
            header_rows.append(row)
    if not data_rows:
        return []

    width = max(len(row) for row in rows)
    labels = _column_labels(header_rows, width)

    rendered_rows: list[str] = []
    group_levels: list[str | None] = []
    for row in data_rows:
        first_number = next(index for index, value in enumerate(row) if _is_number(value))
        # 행 라벨: 숫자 이전의 문자열 셀(계층형). 상위 칸이 비어 있으면 위 행의 값을 이어 받는다.
        levels = [one_line(str(value)) if value not in (None, "") else None for value in row[:first_number]]
        merged: list[str] = []
        for level, value in enumerate(levels):
            if value is not None:
                group_levels = group_levels[:level] + [value]
            elif level < len(group_levels):
                value = group_levels[level]
            if value:
                merged.append(value)
        row_label = " ".join(merged) or "전체"

        cells: list[str] = []
        sample = None
        for column in range(first_number, len(row)):
            value = row[column]
            if not _is_number(value):
                continue
            label = labels[column] if column < len(labels) else ""
            if "가중 사례수" in label and "비가중" not in label:
                continue  # 가중 사례수는 모집단 추정치로 토큰만 차지한다.
            if "사례수" in label or label.upper() == "CASES":
                sample = _fmt_number(value)
                continue
            if value == 0:
                continue
            cells.append(f"{label or '값'} {_fmt_number(value)}")
        if not cells:
            continue
        prefix = f"{row_label}(n={sample})" if sample else row_label
        rendered_rows.append(f"- {prefix}: " + ", ".join(cells))

    if not rendered_rows:
        return []

    header = f"표: {title}\n({note})"
    chunks: list[str] = []
    buffer = header
    for line in rendered_rows:
        if len(buffer) + len(line) + 1 > TABLE_CHUNK_CHARS and buffer != header:
            chunks.append(buffer)
            buffer = f"{header} (계속)"
        buffer = f"{buffer}\n{line}"
    chunks.append(buffer)
    return chunks


def load_stats_xlsx(document: SourceDocument, path: Path) -> Iterator[dict]:
    from openpyxl import load_workbook

    workbook = load_workbook(str(path), read_only=True, data_only=True)
    try:
        # 목차/Index 시트를 제외하고 실제 표가 들어 있는 시트만 읽는다.
        sheets = [name for name in workbook.sheetnames if not re.match(r"(?i)^(index\d*|목차)$", name.strip())]
        note = "값은 원 통계표 수치. 비율 문항은 %, 평균 문항은 해당 단위. 0인 항목은 생략"
        extra = document.options.get("note")
        if extra:
            note = f"{note}. {extra}"

        counter = itertools.count()
        for sheet_name in sheets:
            # 통계표는 '제목 행 → 헤더 행 → 데이터 행 → 빈 행' 블록이 반복되는 구조다.
            block: list[tuple] = []
            for row in workbook[sheet_name].iter_rows(values_only=True):
                if all(value in (None, "") for value in row):
                    if block:
                        yield from _emit_table(document, block, note, counter)
                        block = []
                    continue
                block.append(row)
            if block:
                yield from _emit_table(document, block, note, counter)
    finally:
        workbook.close()


def _emit_table(document: SourceDocument, block: list[tuple], note: str, counter: Iterator[int]) -> Iterator[dict]:
    title_cells = [one_line(str(value)) for value in block[0] if value not in (None, "")]
    if not title_cells or any(_is_number(value) for value in block[0]):
        return
    title = title_cells[0].strip("▣ ").strip() or title_cells[0]
    for piece in _render_table(title, block[1:], note):
        yield _chunk(document, next(counter), f"통계표 「{title}」", piece, {"table": title})


# ---------------------------------------------------------------------------
# 서울 열린데이터광장 JSON
# ---------------------------------------------------------------------------

def _load_open_data(path: Path) -> list[dict]:
    with path.open("r", encoding="utf-8") as file:
        payload = json.load(file)
    return payload.get("DATA", []) if isinstance(payload, dict) else payload


def _field_lines(pairs: list[tuple[str, str | None]], limit: int = 300) -> str:
    lines = []
    for label, value in pairs:
        text = one_line(value, limit)
        if text:
            lines.append(f"{label}: {text}")
    return "\n".join(lines)


def _district_from(address: str) -> str | None:
    match = re.search(r"([가-힣]+구)(?:\s|$)", address or "")
    return match.group(1) if match else None


def load_attractions_json(document: SourceDocument, path: Path) -> Iterator[dict]:
    languages = set(document.options.get("languages", ["ko"]))
    seq = 0
    for item in _load_open_data(path):
        if item.get("lang_code_id") not in languages:
            continue
        name = one_line(item.get("post_sj"))
        if not name:
            continue
        address = one_line(item.get("new_address") or item.get("address"))
        text = _field_lines([
            ("명소", name),
            ("주소", address),
            ("운영시간", item.get("cmmn_use_time")),
            ("운영요일", item.get("cmmn_bsnde")),
            ("휴무일", item.get("cmmn_rstde")),
            ("교통", item.get("subway_info")),
            ("전화", item.get("cmmn_telno")),
            ("태그", item.get("tag")),
            ("장애인 편의시설", item.get("bf_desc")),
            ("웹사이트", item.get("cmmn_hmpg_url")),
        ])
        yield _chunk(document, seq, name, text, {
            "name": name,
            "district": _district_from(address),
            "url": item.get("post_url") or "",
        })
        seq += 1


def load_culture_spaces_json(document: SourceDocument, path: Path) -> Iterator[dict]:
    seq = 0
    for item in _load_open_data(path):
        name = one_line(item.get("fac_name"))
        if not name:
            continue
        text = _field_lines([
            ("문화공간", name),
            ("분류", item.get("subjcode")),
            ("자치구", item.get("gngu")),
            ("주소", item.get("addr")),
            ("관람시간", item.get("openhour")),
            ("휴관일", item.get("closeday")),
            ("관람료", item.get("entr_fee")),
            ("무료구분", item.get("entrfree")),
            ("지하철", item.get("subway")),
            ("버스", item.get("busstop")),
            ("전화", item.get("phne")),
            ("홈페이지", item.get("homepage")),
            ("소개", clean_html(item.get("fac_desc"))),
        ], limit=350)
        yield _chunk(document, seq, name, text, {
            "name": name,
            "district": one_line(item.get("gngu")) or None,
            "latitude": _to_float(item.get("x_coord")),
            "longitude": _to_float(item.get("y_coord")),
            "image_url": item.get("main_img") or "",
            "url": item.get("homepage") or "",
        })
        seq += 1


def load_culture_events_json(document: SourceDocument, path: Path, today: date | None = None) -> Iterator[dict]:
    today = today or date.today()
    cutoff = today - timedelta(days=int(document.options.get("ended_within_days", 0)))
    seq = 0
    skipped = 0
    for item in _load_open_data(path):
        title = one_line(item.get("title"))
        start, end = _parse_period(item.get("date"))
        if not title or not end:
            continue
        if end < cutoff.isoformat():
            skipped += 1
            continue
        text = _field_lines([
            ("행사", title),
            ("분류", item.get("codename")),
            ("기간", f"{start} ~ {end}"),
            ("시간", item.get("pro_time")),
            ("자치구", item.get("guname")),
            ("장소", item.get("place")),
            ("대상", item.get("use_trgt")),
            ("요금", item.get("use_fee")),
            ("유무료", item.get("is_free")),
            ("주최", item.get("org_name")),
            ("출연", item.get("player")),
            ("소개", item.get("program")),
            ("상세", item.get("hmpg_addr")),
        ], limit=250)
        yield _chunk(document, seq, title, text, {
            "name": title,
            "category": one_line(item.get("codename")),
            "district": one_line(item.get("guname")) or None,
            "start_date": start,
            "end_date": end,
            "latitude": _to_float(item.get("lat")),
            "longitude": _to_float(item.get("lot")),
            "image_url": item.get("main_img") or "",
            "url": item.get("hmpg_addr") or item.get("org_link") or "",
        })
        seq += 1
    logger.info("문화행사: %s건 색인, 종료된 행사 %s건 제외(기준일 %s)", seq, skipped, cutoff)


def _to_float(value) -> float | None:
    try:
        return float(value) if value not in (None, "") else None
    except (TypeError, ValueError):
        return None


# ---------------------------------------------------------------------------
# 자치구 CSV / 관광지 분석 CSV
# ---------------------------------------------------------------------------

DISTRICT_FILE_PATTERN = re.compile(r"서울특별시_([가-힣]+구)_([가-힣]+)")
DISTRICT_NAME_FIELDS = ("길명", "관광지명", "거리명")
DISTRICT_SKIP_FIELDS = {"관리기관전화번호"}


def load_district_csv(document: SourceDocument, path: Path) -> Iterator[dict]:
    match = DISTRICT_FILE_PATTERN.search(path.stem)
    district, kind = (match.group(1), match.group(2)) if match else (None, path.stem)
    for index, row in enumerate(_read_csv(path)):
        name = next((row[field] for field in DISTRICT_NAME_FIELDS if row.get(field)), "")
        if not name:
            continue
        pairs = [("유형", kind), ("자치구", district)]
        pairs += [(key, value) for key, value in row.items() if key and key not in DISTRICT_SKIP_FIELDS]
        text = _field_lines(pairs, limit=400)
        yield {
            **_chunk(document, 0, f"{district or ''} {kind} · {name}".strip(), text, {
                "name": name,
                "district": district,
                "kind": kind,
                "latitude": _to_float(row.get("위도")),
                "longitude": _to_float(row.get("경도")),
            }),
            "id": f"{document.id}:{path.stem}:{index:04d}",
        }


ANALYTICS_HEADER_PATTERN = re.compile(r"\([A-Z0-9_]+\)")
ANALYTICS_ROW_LIMIT = 40


def load_analytics_csv(document: SourceDocument, path: Path) -> Iterator[dict]:
    rows = _read_csv(path)
    if not rows:
        return
    dataset = path.stem.replace("서울 관광지_", "").replace("_", " ")
    name_key = next((key for key in rows[0] if key.startswith("서울관광지_명")), None)
    grouped: dict[str, list[dict]] = defaultdict(list)
    for row in rows:
        grouped[row.get(name_key, "") if name_key else ""].append(row)

    for index, (place, items) in enumerate(sorted(grouped.items())):
        lines = [f"데이터: 서울 관광지 {dataset}", f"관광지: {place or '전체'}"]
        for row in items[:ANALYTICS_ROW_LIMIT]:
            cells = []
            for key, value in row.items():
                if not value or key == name_key or "_코드" in key or "형상정보" in key:
                    continue
                cells.append(f"{ANALYTICS_HEADER_PATTERN.sub('', key).strip()} {value}")
            lines.append("- " + ", ".join(cells))
        if len(items) > ANALYTICS_ROW_LIMIT:
            lines.append(f"(이하 {len(items) - ANALYTICS_ROW_LIMIT}행 생략)")
        yield {
            **_chunk(document, 0, f"{place} · {dataset}", "\n".join(lines), {"name": place, "dataset": dataset}),
            "id": f"{document.id}:{path.stem}:{index:04d}",
        }


LOADERS = {
    "pdf": load_pdf,
    "stats_xlsx": load_stats_xlsx,
    "attractions_json": load_attractions_json,
    "culture_spaces_json": load_culture_spaces_json,
    "culture_events_json": load_culture_events_json,
    "district_csv": load_district_csv,
    "analytics_csv": load_analytics_csv,
}


def load_document(document: SourceDocument, data_dir: Path, today: date | None = None) -> Iterator[dict]:
    loader = LOADERS.get(document.loader)
    if loader is None:
        raise ValueError(f"지원하지 않는 loader입니다: {document.loader} ({document.id})")
    paths = document.resolve(data_dir)
    if not paths:
        logger.warning("원천 파일이 없어 건너뜁니다: %s (%s)", document.id, document.path)
        return
    for path in paths:
        if loader is load_culture_events_json:
            yield from loader(document, path, today=today)
        else:
            yield from loader(document, path)
